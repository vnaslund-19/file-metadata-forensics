"""Metadata stored inside Office documents, read through the format libraries.

Reading the same properties straight out of the file is done separately, in
ooxml.py, so the two readings can be compared rather than sharing code.

The libraries only expose the core properties. Company, Manager and Hyperlink
base live in docProps/app.xml, which none of them read.
"""

import datetime

import docx
import openpyxl
import pptx

# Every result has all of these. Fields a format doesn't have are None.
METADATA_FIELDS = [
    "title",
    "subject",
    "author",
    "keywords",
    "comments",
    "category",
    "last_modified_by",
    "revision",
    "created",
    "modified",
    "last_printed",
    "application",
    "app_version",
    "company",
    "manager",
    "hyperlink_base",
    "template",
    "producer",
    "pages",
    "pdf_version",
    "has_xmp_metadata",
    "document_id",
    "instance_id",
    "content_status",
    "identifier",
    "language",
    "version",
]

# Core property names shared by DOCX and PPTX, as python-docx and python-pptx
# spell them.
OOXML_CORE_PROPERTIES = [
    "title",
    "subject",
    "author",
    "keywords",
    "comments",
    "category",
    "content_status",
    "identifier",
    "language",
    "version",
    "last_modified_by",
    "revision",
    "created",
    "modified",
    "last_printed",
]

# openpyxl spells several of the same properties differently.
XLSX_CORE_PROPERTIES = {
    "title": "title",
    "subject": "subject",
    "author": "creator",
    "keywords": "keywords",
    "comments": "description",
    "category": "category",
    "content_status": "contentStatus",
    "identifier": "identifier",
    "language": "language",
    "version": "version",
    "last_modified_by": "lastModifiedBy",
    "revision": "revision",
    "created": "created",
    "modified": "modified",
    "last_printed": "lastPrinted",
}


def extract(path, file_type):
    """Return (metadata, warnings). A file that can't be read gives a warning, not a crash."""
    extractors = {
        "docx": extract_docx,
        "xlsx": extract_xlsx,
        "pptx": extract_pptx,
    }
    metadata = dict.fromkeys(METADATA_FIELDS)
    try:
        found, warnings = extractors[file_type](path)
    except Exception as error:
        return metadata, [f"could not read {file_type} metadata: {error}"]
    metadata.update(found)
    return metadata, warnings


def extract_docx(path):
    return read_ooxml_core(docx.Document(path).core_properties)


def extract_pptx(path):
    return read_ooxml_core(pptx.Presentation(path).core_properties)


def read_ooxml_core(props):
    metadata = {}
    for name in OOXML_CORE_PROPERTIES:
        metadata[name] = clean(getattr(props, name, None))
    return metadata, []


def extract_xlsx(path):
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        props = workbook.properties
        metadata = {}
        for name, source_name in XLSX_CORE_PROPERTIES.items():
            metadata[name] = clean(getattr(props, source_name, None))
        # openpyxl gives the revision as text, python-docx and python-pptx as a number.
        if metadata["revision"] and metadata["revision"].isdigit():
            metadata["revision"] = int(metadata["revision"])
        return metadata, []
    finally:
        workbook.close()


def clean(value):
    """Normalise a property value for output."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        # Office core properties are stored as UTC, but openpyxl hands back a
        # datetime with no timezone on it. Say so rather than leaving dates
        # from different formats looking inconsistent.
        if value.tzinfo is None:
            value = value.replace(tzinfo=datetime.timezone.utc)
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value
