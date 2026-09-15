"""Create sample files with known metadata using Python libraries.

These files are written by python-docx, openpyxl, python-pptx and pypdf instead
of Microsoft Office, so their metadata does not look the same as the
Office-authored files in samples/. They give the extraction code something to
run against that can be recreated at any time, on a machine without Office.

Usage:
    python tests/make_generated_samples.py
"""

import datetime
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdf

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "samples" / "generated"

CREATED = datetime.datetime(2026, 1, 2, 3, 4, 5)
MODIFIED = datetime.datetime(2026, 1, 2, 6, 7, 8)

TITLE = "Generated Test Document"
SUBJECT = "Generated Subject"
AUTHOR = "Generated Author"
KEYWORDS = "generated, metadata, test"
COMMENTS = "Written by tests/make_generated_samples.py"
CATEGORY = "Generated"
LAST_MODIFIED_BY = "Generated Editor"
HYPERLINK = "https://example.com/generated-link-target"


def set_core_properties(props):
    props.title = TITLE
    props.subject = SUBJECT
    props.author = AUTHOR
    props.keywords = KEYWORDS
    props.comments = COMMENTS
    props.category = CATEGORY
    props.last_modified_by = LAST_MODIFIED_BY
    props.revision = 2
    props.created = CREATED
    props.modified = MODIFIED


def add_hyperlink(paragraph, url, text):
    """python-docx has no hyperlink API, so build the link and relationship."""
    rel_id = paragraph.part.relate_to(
        url,
        docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    link = docx.oxml.shared.OxmlElement("w:hyperlink")
    link.set(docx.oxml.ns.qn("r:id"), rel_id)
    run = docx.oxml.shared.OxmlElement("w:r")
    run_text = docx.oxml.shared.OxmlElement("w:t")
    run_text.text = text
    run.append(run_text)
    link.append(run)
    paragraph._p.append(link)


def make_docx(path):
    document = docx.Document()
    document.add_paragraph("Generated DOCX body text.")
    add_hyperlink(document.add_paragraph(), HYPERLINK, "generated link")
    set_core_properties(document.core_properties)
    document.save(path)


def make_xlsx(path):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Generated XLSX cell."
    props = workbook.properties
    props.title = TITLE
    props.subject = SUBJECT
    props.creator = AUTHOR
    props.keywords = KEYWORDS
    props.description = COMMENTS
    props.category = CATEGORY
    props.lastModifiedBy = LAST_MODIFIED_BY
    props.revision = "2"
    props.created = CREATED
    # openpyxl replaces the modified time with the save time, so it is not set
    # here and cannot be relied on as a known value.
    workbook.save(path)


def make_pptx(path):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Generated PPTX slide."
    set_core_properties(presentation.core_properties)
    presentation.save(path)


def make_pdf(path):
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_metadata(
        {
            "/Title": TITLE,
            "/Subject": SUBJECT,
            "/Author": AUTHOR,
            "/Keywords": KEYWORDS,
            "/Creator": "tests/make_generated_samples.py",
            "/Producer": "pypdf",
            "/CreationDate": "D:20260102030405Z",
            "/ModDate": "D:20260102060708Z",
        }
    )
    with open(path, "wb") as handle:
        writer.write(handle)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    make_docx(OUTPUT_DIR / "generated.docx")
    make_xlsx(OUTPUT_DIR / "generated.xlsx")
    make_pptx(OUTPUT_DIR / "generated.pptx")
    make_pdf(OUTPUT_DIR / "generated.pdf")
    for path in sorted(OUTPUT_DIR.iterdir()):
        print(f"wrote {path.relative_to(OUTPUT_DIR.parent.parent)}")


if __name__ == "__main__":
    main()
